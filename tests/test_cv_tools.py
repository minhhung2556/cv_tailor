import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parents[1]
SKILLS = REPO_ROOT / "skills" / "cv-tailor"
AGENT_SKILLS = REPO_ROOT / ".agents" / "skills" / "cv-tailor"
SCRIPTS = SKILLS / "scripts"
HARVARD = SKILLS / "resources" / "harvard_cv.typ"


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _stage_installed_skill(
    base: Path,
    *,
    skill_parts: tuple[str, ...],
    script_names: tuple[str, ...],
) -> Path:
    skill = base.joinpath(*skill_parts)
    (skill / "resources").mkdir(parents=True)
    (skill / "resources" / "harvard_cv.typ").write_text(
        HARVARD.read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    for name in script_names:
        dest = skill / "scripts" / name
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes((SCRIPTS / name).read_bytes())
    return skill


parse_jd = _load_module("parse_jd", SCRIPTS / "parse_jd.py")
extract_cv = _load_module("extract_cv", SCRIPTS / "extract_cv.py")
compile_cv = _load_module("compile_cv", SCRIPTS / "compile_cv.py")
link_agent_skill = _load_module("link_agent_skill", SCRIPTS / "link_agent_skill.py")
cv_io = _load_module("cv_io", SCRIPTS / "cv_io.py")
build_docx_cv = _load_module("build_docx_cv", SCRIPTS / "build_docx_cv.py")
docx_to_pdf = _load_module("docx_to_pdf", SCRIPTS / "docx_to_pdf.py")


def _write_minimal_docx(path: Path, paragraph: str) -> None:
    try:
        from docx import Document
    except ImportError:
        raise unittest.SkipTest("python-docx not installed")
    doc = Document()
    doc.add_paragraph(paragraph)
    doc.save(str(path))


class ParseJdTests(unittest.TestCase):
    def test_html_strips_chrome_and_normalizes_whitespace(self):
        html = """<!DOCTYPE html>
<html><head><title>JD</title><style>.x{}</style></head>
<body>
<nav>Skip to content</nav>
<header>Senior Staff Engineer</header>
<main><p>Senior  Engineer</p><script>alert(1)</script></main>
<footer>Copyright</footer>
</body></html>"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.html"
            path.write_text(html, encoding="utf-8")
            text = parse_jd.parse_jd(str(path))

        self.assertIn("Senior Engineer", text)
        self.assertIn("Senior Staff Engineer", text)
        self.assertNotIn("Skip to content", text)
        self.assertNotIn("Copyright", text)
        self.assertNotIn("alert", text)

    def test_html_keeps_title_when_not_in_body(self):
        html = """<!DOCTYPE html>
<html><head><title>Senior Backend Engineer - Acme</title></head>
<body><main><p>About the role</p></main></body></html>"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.html"
            path.write_text(html, encoding="utf-8")
            text = parse_jd.parse_jd(str(path))
        self.assertIn("Senior Backend Engineer - Acme", text)
        self.assertIn("About the role", text)

    def test_cli_json_flag_before_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.txt"
            path.write_text("hello", encoding="utf-8")
            buf = io.StringIO()
            with mock.patch.object(sys, "argv", ["parse_jd.py", "--json", str(path)]):
                with mock.patch.object(sys, "stdout", buf):
                    parse_jd.main()
        self.assertIn('"text": "hello"', buf.getvalue())

    def test_missing_file_raises(self):
        with self.assertRaises(FileNotFoundError):
            parse_jd.parse_jd("nonexistent_jd.html")

    def test_unsupported_extension_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.xyz"
            path.write_bytes(b"fake")
            with self.assertRaises(ValueError):
                parse_jd.parse_jd(str(path))

    def test_legacy_doc_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.doc"
            path.write_bytes(b"\xd0\xcf\x11\xe0")
            with self.assertRaises(RuntimeError) as ctx:
                parse_jd.parse_jd(str(path))
            self.assertIn(".doc", str(ctx.exception))
            self.assertIn("not supported", str(ctx.exception).lower())

    def test_docx_extracts_paragraph_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.docx"
            _write_minimal_docx(path, "Senior Backend Engineer role")
            text = parse_jd.parse_jd(str(path))
        self.assertIn("Senior Backend Engineer role", text)

    def test_docx_cli_json_shape(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.docx"
            _write_minimal_docx(path, "hello from docx")
            buf = io.StringIO()
            with mock.patch.object(sys, "argv", ["parse_jd.py", str(path), "--json"]):
                with mock.patch.object(sys, "stdout", buf):
                    parse_jd.main()
        payload = json.loads(buf.getvalue())
        self.assertEqual(payload["file_name"], "jd.docx")
        self.assertIn("hello from docx", payload["text"])

    def test_markdown_collapses_whitespace(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.md"
            path.write_text("Role   title\n\nSecond  line\n", encoding="utf-8")
            text = parse_jd.parse_jd(str(path))
        self.assertEqual(text, "Role title\nSecond line")

    def test_docx_table_dedupes_merged_cell_text(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jd.docx"
            doc = Document()
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Same requirement"
            table.rows[0].cells[1].text = "Same requirement"
            doc.save(str(path))
            text = parse_jd.parse_jd(str(path))
        self.assertEqual(text.count("Same requirement"), 1)


class BuildDocxCvTests(unittest.TestCase):
    def test_build_writes_name_and_bullet(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx not installed")

        spec = {
            "name": "Jane Candidate",
            "title": "Engineer",
            "contact": "jane@example.com",
            "sections": [
                {
                    "heading": "Core Competencies",
                    "type": "bullets",
                    "items": ["Shipped the payment API"],
                },
                {
                    "heading": "Experience",
                    "type": "experience",
                    "entries": [
                        {
                            "company": "Acme",
                            "location": "Remote",
                            "role": "Senior Engineer",
                            "dates": "2020 – Present",
                            "bullets": ["Led the platform migration"],
                        }
                    ],
                },
            ],
        }
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "cv.docx"
            build_docx_cv.build(spec, str(out))
            self.assertTrue(out.is_file())
            doc = Document(str(out))
            full_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("JANE CANDIDATE", full_text)
            self.assertIn("Shipped the payment API", full_text)


class DocxToPdfTests(unittest.TestCase):
    def test_no_converter_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "in.docx"
            _write_minimal_docx(docx_path, "x")
            pdf_path = Path(tmp) / "out.pdf"
            with mock.patch.object(docx_to_pdf, "_convert_with_word", return_value=False):
                with mock.patch.object(docx_to_pdf, "_convert_with_libreoffice", return_value=False):
                    with self.assertRaises(RuntimeError) as ctx:
                        docx_to_pdf.docx_to_pdf(str(docx_path), str(pdf_path))
            self.assertIn("converter", str(ctx.exception).lower())

    def test_word_failure_still_tries_libreoffice(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "in.docx"
            _write_minimal_docx(docx_path, "x")
            pdf_path = Path(tmp) / "out.pdf"
            lo = mock.Mock(return_value=False)

            def word_broken(*_args, **_kwargs):
                raise OSError("COM failed")

            with mock.patch.object(docx_to_pdf, "_convert_with_word", side_effect=word_broken):
                with mock.patch.object(docx_to_pdf, "_convert_with_libreoffice", lo):
                    with self.assertRaises(RuntimeError):
                        docx_to_pdf.docx_to_pdf(str(docx_path), str(pdf_path))
            lo.assert_called_once()

    def test_stale_destination_pdf_not_returned_on_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "in.docx"
            _write_minimal_docx(docx_path, "x")
            pdf_path = Path(tmp) / "out.pdf"
            pdf_path.write_bytes(b"%PDF-old\n")
            with mock.patch.object(docx_to_pdf, "_convert_with_word", return_value=False):
                with mock.patch.object(docx_to_pdf, "_convert_with_libreoffice", return_value=False):
                    with self.assertRaises(RuntimeError):
                        docx_to_pdf.docx_to_pdf(str(docx_path), str(pdf_path))
            self.assertEqual(pdf_path.read_bytes(), b"%PDF-old\n")


class ExtractCvTests(unittest.TestCase):
    def test_typ_returns_source_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cv.typ"
            path.write_text("#let x = 1\n", encoding="utf-8")
            text = extract_cv.extract_cv(str(path))
            self.assertEqual(text, "#let x = 1\n")

    def test_prefers_sibling_typ_over_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp) / "cv_base"
            base.with_suffix(".pdf").write_bytes(b"%PDF-1.4\n")
            base.with_suffix(".typ").write_text("from typ\n", encoding="utf-8")
            text = extract_cv.extract_cv(str(base.with_suffix(".pdf")))
            self.assertEqual(text, "from typ\n")

    def test_cli_json_flag_before_default_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            (workspace / "assets").mkdir()
            (workspace / "assets" / "cv_base.typ").write_text("from default\n", encoding="utf-8")
            buf = io.StringIO()
            with mock.patch.object(extract_cv.cv_io, "find_workspace_root", return_value=workspace):
                with mock.patch.object(sys, "argv", ["extract_cv.py", "--json"]):
                    with mock.patch.object(sys, "stdout", buf):
                        extract_cv.main()
            out = buf.getvalue()
            self.assertIn('"text":', out)
            self.assertIn("from default", out)

    def test_missing_file_raises(self):
        with self.assertRaises(FileNotFoundError):
            extract_cv.extract_cv("nonexistent_cv.typ")


class CvIoPdfTests(unittest.TestCase):
    def test_empty_pdf_raises(self):
        try:
            import pypdf
        except ImportError:
            self.skipTest("pypdf not installed")

        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "blank.pdf"
            with open(path, "wb") as f:
                writer.write(f)
            with self.assertRaises(RuntimeError):
                cv_io.extract_pdf_text(path)


class CompileCvTests(unittest.TestCase):
    def test_find_skill_root_from_script_path(self):
        root = cv_io.find_skill_root(SCRIPTS)
        self.assertEqual(root, SKILLS)

    def test_find_workspace_root_from_repo(self):
        root = cv_io.find_workspace_root(REPO_ROOT)
        self.assertEqual(root, REPO_ROOT)

    def test_find_workspace_root_falls_back_to_skills_grandparent(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            repo = base / "repo"
            scripts = repo / "skills" / "cv-tailor" / "scripts"
            scripts.mkdir(parents=True)
            (repo / "skills" / "cv-tailor" / "resources").mkdir(parents=True)
            (repo / "skills" / "cv-tailor" / "resources" / "harvard_cv.typ").write_text(
                HARVARD.read_text(encoding="utf-8"),
                encoding="utf-8",
            )
            (scripts / "cv_io.py").write_bytes((SCRIPTS / "cv_io.py").read_bytes())
            iso_cv_io = _load_module("iso_cv_io_fallback", scripts / "cv_io.py")
            cwd = base / "nowhere"
            cwd.mkdir()
            root = iso_cv_io.find_workspace_root(cwd)
            self.assertEqual(root, repo.resolve())

    def test_resolve_relative_paths_from_workspace(self):
        workspace = cv_io.find_workspace_root(REPO_ROOT)
        typ = cv_io.resolve_workspace_path(workspace, "assets/cv_base.typ")
        pdf = cv_io.resolve_workspace_path(workspace, "output/smoke.pdf")
        self.assertEqual(typ, (workspace / "assets/cv_base.typ").resolve())
        self.assertEqual(pdf, (workspace / "output/smoke.pdf").resolve())

    def test_stage_harvard_template_next_to_typ(self):
        with tempfile.TemporaryDirectory() as tmp:
            typ = Path(tmp) / "output" / "tuned.typ"
            typ.parent.mkdir(parents=True)
            typ.write_text("#import \"harvard_cv.typ\"\n", encoding="utf-8")
            dest = cv_io.stage_harvard_template_next_to(typ, SKILLS)
            self.assertEqual(dest, typ.parent / "harvard_cv.typ")
            self.assertTrue(dest.is_file())
            self.assertIn("#let cv_bullet", dest.read_text(encoding="utf-8"))

    def test_compile_cv_default_pdf_beside_typ(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp) / "workspace"
            (workspace / "assets").mkdir(parents=True)
            (workspace / "assets" / "cv_base.typ").write_text("x", encoding="utf-8")
            typ = workspace / "assets" / "tuned.typ"
            typ.write_text("#import \"harvard_cv.typ\"\n", encoding="utf-8")
            expected_pdf = typ.with_suffix(".pdf")
            expected_pdf.write_bytes(b"%PDF-1.4\n")

            def fake_run(cmd, **kwargs):
                self.assertEqual(cmd[-1], str(expected_pdf))
                return mock.Mock(returncode=0, stdout="", stderr="")

            with mock.patch.object(compile_cv.cv_io, "find_workspace_root", return_value=workspace):
                with mock.patch.object(compile_cv.cv_io, "stage_harvard_template_next_to"):
                    with mock.patch.object(compile_cv.subprocess, "run", side_effect=fake_run):
                        result = compile_cv.compile_cv(str(typ))
            self.assertEqual(result, str(expected_pdf))

    def test_compile_cv_typst_root_when_typ_outside_workspace(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            workspace = base / "workspace"
            (workspace / "assets").mkdir(parents=True)
            (workspace / "assets" / "cv_base.typ").write_text("x", encoding="utf-8")
            outside = base / "outside"
            outside.mkdir()
            typ = outside / "tuned.typ"
            typ.write_text("#import \"harvard_cv.typ\"\n", encoding="utf-8")
            expected_pdf = typ.with_suffix(".pdf")
            expected_pdf.write_bytes(b"%PDF-1.4\n")
            captured_root: list[str] = []

            def fake_run(cmd, **kwargs):
                captured_root.append(cmd[cmd.index("--root") + 1])
                return mock.Mock(returncode=0, stdout="", stderr="")

            with mock.patch.object(compile_cv.cv_io, "find_workspace_root", return_value=workspace):
                with mock.patch.object(compile_cv.cv_io, "stage_harvard_template_next_to"):
                    with mock.patch.object(compile_cv.subprocess, "run", side_effect=fake_run):
                        compile_cv.compile_cv(str(typ))
            self.assertEqual(captured_root[0], str(outside.resolve()))


class AgentSkillLinkTests(unittest.TestCase):
    def _minimal_repo(self, tmp: Path, *, extra_skills: tuple[str, ...] = ()) -> Path:
        root = tmp / "repo"
        (root / "skills" / "cv-tailor" / "resources").mkdir(parents=True)
        (root / "skills" / "cv-tailor" / "resources" / "harvard_cv.typ").write_text(
            HARVARD.read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        (root / "skills" / "cv-tailor" / "scripts").mkdir(parents=True, exist_ok=True)
        (root / "skills" / "cv-tailor" / "SKILL.md").write_text(
            "---\nname: cv-tailor\n---\n",
            encoding="utf-8",
        )
        for name in extra_skills:
            skill_dir = root / "skills" / name
            skill_dir.mkdir(parents=True, exist_ok=True)
            (skill_dir / "SKILL.md").write_text(f"---\nname: {name}\n---\n", encoding="utf-8")
        return root

    def test_real_directory_without_force_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = self._minimal_repo(Path(tmp))
            agent = root / ".agents" / "skills" / "cv-tailor"
            agent.mkdir(parents=True)
            (agent / "keep.txt").write_text("data", encoding="utf-8")
            with self.assertRaises(RuntimeError):
                link_agent_skill.ensure_agent_skill_junction(root, force=False)

    def test_junction_points_at_skills_tree_in_temp_repo(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = self._minimal_repo(Path(tmp))
            link_agent_skill.ensure_agent_skill_junction(root)
            skills = (root / "skills" / "cv-tailor").resolve()
            for rel in (".agents/skills/cv-tailor", ".claude/skills/cv-tailor"):
                agent = root / rel
                self.assertTrue(agent.exists())
                self.assertEqual(agent.resolve(), skills)

    def test_reparse_point_replaced(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = self._minimal_repo(Path(tmp))
            agent = root / ".agents" / "skills" / "cv-tailor"
            agent.parent.mkdir(parents=True, exist_ok=True)
            skills = (root / "skills" / "cv-tailor").resolve()
            if sys.platform == "win32":
                import subprocess

                subprocess.run(
                    ["cmd", "/c", "mklink", "/J", str(agent), str(skills)],
                    check=True,
                    capture_output=True,
                )
            else:
                agent.symlink_to(skills, target_is_directory=True)
            marker = root / "skills" / "cv-tailor" / "resources" / "marker.txt"
            marker.write_text("ok", encoding="utf-8")
            link_agent_skill.ensure_agent_skill_junction(root)
            self.assertTrue((agent / "resources" / "marker.txt").is_file())

    def test_extra_skill_gets_both_agent_links(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = self._minimal_repo(Path(tmp), extra_skills=("github-contributions",))
            link_agent_skill.ensure_agent_skill_junction(root)
            gh_source = (root / "skills" / "github-contributions").resolve()
            for rel in (
                ".agents/skills/github-contributions",
                ".claude/skills/github-contributions",
            ):
                agent = root / rel
                self.assertTrue(agent.exists())
                self.assertEqual(agent.resolve(), gh_source)

class InstalledSkillWorkspaceTests(unittest.TestCase):
    def test_parse_jd_resolves_relative_path_from_separate_workspace(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            skill = _stage_installed_skill(
                base,
                skill_parts=("installed", "cv-tailor"),
                script_names=("parse_jd.py", "extract_cv.py", "compile_cv.py", "cv_io.py"),
            )

            workspace = base / "my-project"
            (workspace / "assets" / "jd").mkdir(parents=True)
            (workspace / "assets" / "cv_base.typ").write_text("base\n", encoding="utf-8")
            jd = workspace / "assets" / "jd" / "role.txt"
            jd.write_text("Backend engineer", encoding="utf-8")

            iso_parse = _load_module("iso_parse_jd", skill / "scripts" / "parse_jd.py")

            with mock.patch.object(iso_parse.cv_io, "find_workspace_root", return_value=workspace):
                text = iso_parse.parse_jd("assets/jd/role.txt")
            self.assertEqual(text, "Backend engineer")

    def test_extract_default_cv_in_separate_workspace(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            skill = _stage_installed_skill(
                base,
                skill_parts=("cv-tailor",),
                script_names=("extract_cv.py", "cv_io.py"),
            )
            (skill / "resources" / "harvard_cv.typ").write_text("x", encoding="utf-8")

            workspace = base / "workspace"
            (workspace / "assets").mkdir(parents=True)
            (workspace / "assets" / "cv_base.typ").write_text("workspace cv\n", encoding="utf-8")

            iso_extract = _load_module("iso_extract", skill / "scripts" / "extract_cv.py")

            with mock.patch.object(iso_extract.cv_io, "find_workspace_root", return_value=workspace):
                default = iso_extract.default_cv_path()
                text = iso_extract.extract_cv(default)
            self.assertIn("cv_base.typ", default)
            self.assertEqual(text, "workspace cv\n")


class CursorPluginManifestTests(unittest.TestCase):
    PLUGIN_JSON = REPO_ROOT / ".cursor-plugin" / "plugin.json"
    MARKETPLACE_JSON = REPO_ROOT / ".cursor-plugin" / "marketplace.json"
    SKILL_MD = SKILLS / "SKILL.md"

    def test_plugin_and_marketplace_manifests(self):
        plugin = json.loads(self.PLUGIN_JSON.read_text(encoding="utf-8"))
        marketplace = json.loads(self.MARKETPLACE_JSON.read_text(encoding="utf-8"))

        self.assertEqual(plugin["name"], "cv-tailor")
        self.assertIn("description", plugin)
        self.assertEqual(plugin.get("skills"), "skills")

        self.assertEqual(marketplace["name"], "cv-tailor")
        self.assertGreaterEqual(len(marketplace["plugins"]), 1)
        entry = marketplace["plugins"][0]
        self.assertEqual(entry["name"], plugin["name"])
        self.assertEqual(entry["source"], ".")

        skill_text = self.SKILL_MD.read_text(encoding="utf-8")
        self.assertIn("name: cv-tailor", skill_text.split("---", 2)[1])


class ClaudePluginManifestTests(unittest.TestCase):
    PLUGIN_JSON = REPO_ROOT / ".claude-plugin" / "plugin.json"
    MARKETPLACE_JSON = REPO_ROOT / ".claude-plugin" / "marketplace.json"
    SKILL_MD = SKILLS / "SKILL.md"

    def test_plugin_and_marketplace_manifests(self):
        plugin = json.loads(self.PLUGIN_JSON.read_text(encoding="utf-8"))
        marketplace = json.loads(self.MARKETPLACE_JSON.read_text(encoding="utf-8"))

        self.assertEqual(plugin["name"], "cv-tailor")
        self.assertIn("description", plugin)

        self.assertEqual(marketplace["name"], "cv-tailor")
        self.assertIn("name", marketplace["owner"])
        self.assertGreaterEqual(len(marketplace["plugins"]), 1)
        entry = marketplace["plugins"][0]
        self.assertEqual(entry["name"], plugin["name"])
        self.assertEqual(entry["source"], ".")

        skill_text = self.SKILL_MD.read_text(encoding="utf-8")
        self.assertIn("claude-code", skill_text)
        self.assertIn("/plugin marketplace add", skill_text)


if __name__ == "__main__":
    unittest.main()
