#!/usr/bin/env python3
"""
build_docx_cv.py - Renders a Harvard-style CV as a Word (.docx) document from a
JSON spec. This is the Word counterpart of resources/harvard_cv.typ: the JSON
schema mirrors the Typst helpers (cv / cv_section / cv_experience / cv_entry) so
the same tuned content can target either PDF (Typst) or Word (this script).

Usage:
    python build_docx_cv.py <spec.json> <output.docx>

JSON schema (see resources/cv_spec.example.json for a full example):
{
  "name": "LUONG DO MINH HUNG",
  "title": "<target role title>",
  "contact": "email · phone · location",
  "sections": [
    {"heading": "Professional Summary", "type": "text",   "content": "..."},
    {"heading": "Core Competencies",    "type": "bullets", "items": ["...", "..."]},
    {"heading": "Experience", "type": "experience", "entries": [
        {"company": "...", "location": "...", "role": "...", "dates": "...",
         "bullets": ["...", "..."]}
    ]},
    {"heading": "Education", "type": "entries", "entries": [
        {"title": "...", "subtitle": "...", "dates": "...", "details": ["..."]}
    ]}
  ]
}
"""

import sys
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --- Layout constants (kept in sync with resources/harvard_cv.typ) ---------
FONT_NAME = "Times New Roman"
BASE_SIZE = 9.5           # pt, body text
NAME_SIZE = 17.0          # pt, header name
TITLE_SIZE = 10.5         # pt, header role title
CONTACT_SIZE = 8.5        # pt, header contact line
SECTION_SIZE = 10.0       # pt, section headings

PAGE_WIDTH = Cm(21.0)     # A4
PAGE_HEIGHT = Cm(29.7)
MARGIN_X = Cm(1.5)
MARGIN_Y = Cm(1.4)
CONTENT_WIDTH = Cm(21.0 - 2 * 1.5)  # right tab stop for dates column

BODY_COLOR = RGBColor(0x11, 0x11, 0x11)
RULE_COLOR = "333333"     # section underline (hex, no '#')


# --- Low-level formatting helpers ------------------------------------------
def _style_run(run, *, size=BASE_SIZE, bold=False, italic=False,
               color=BODY_COLOR, upper=False):
    if upper:
        run.text = run.text.upper()
    run.font.name = FONT_NAME
    # Force the font for all script ranges, otherwise Word may fall back.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(attr), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def _new_paragraph(doc, *, align=None, space_before=0.0, space_after=2.0,
                   justify=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align is not None:
        p.alignment = align
    return p


def _add_bottom_border(paragraph):
    """Draw a thin rule under a paragraph, like the Typst cv_section line."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')       # 6 eighths of a point ≈ 0.75pt
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), RULE_COLOR)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _two_column_line(doc, left_runs, right_text, *, space_before=3.0,
                     space_after=1.0):
    """A row with left content and a right-aligned dates column, mirroring the
    Typst grid(columns: (1fr, auto)). left_runs is a list of (text, kwargs)."""
    p = _new_paragraph(doc, space_before=space_before, space_after=space_after)
    p.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    for text, kwargs in left_runs:
        _style_run(p.add_run(text), **kwargs)
    if right_text:
        _style_run(p.add_run('\t' + right_text), size=BASE_SIZE)
    return p


def _add_bullets(doc, items):
    for item in items:
        p = _new_paragraph(doc, space_before=0.0, space_after=1.0)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.5)
        pf.first_line_indent = Cm(-0.35)  # hanging indent for the bullet glyph
        _style_run(p.add_run('•\t'))
        _style_run(p.add_run(str(item)))


# --- Section renderers (dispatched by "type") ------------------------------
def _render_text(doc, section):
    p = _new_paragraph(doc, justify=True, space_after=2.0)
    _style_run(p.add_run(section.get('content', '')))


def _render_bullets(doc, section):
    _add_bullets(doc, section.get('items', []))


def _render_experience(doc, section):
    for entry in section.get('entries', []):
        left = [(entry.get('company', ''), {'bold': True})]
        if entry.get('location'):
            left.append((', ' + entry['location'], {}))
        _two_column_line(doc, left, entry.get('dates', ''))
        if entry.get('role'):
            p = _new_paragraph(doc, space_before=0.0, space_after=1.0)
            _style_run(p.add_run(entry['role']), italic=True)
        _add_bullets(doc, entry.get('bullets', []))


def _render_entries(doc, section):
    for entry in section.get('entries', []):
        left = [(entry.get('title', ''), {'bold': True})]
        if entry.get('subtitle'):
            left.append((' · ', {}))
            left.append((entry['subtitle'], {'italic': True}))
        _two_column_line(doc, left, entry.get('dates', ''),
                         space_before=2.0)
        _add_bullets(doc, entry.get('details', []))


SECTION_RENDERERS = {
    'text': _render_text,
    'bullets': _render_bullets,
    'experience': _render_experience,
    'entries': _render_entries,
}


# --- Document assembly ------------------------------------------------------
def _configure_document(doc):
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.left_margin = MARGIN_X
    section.right_margin = MARGIN_X
    section.top_margin = MARGIN_Y
    section.bottom_margin = MARGIN_Y

    normal = doc.styles['Normal']
    normal.font.name = FONT_NAME
    normal.font.size = Pt(BASE_SIZE)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(attr), FONT_NAME)


def _add_header(doc, spec):
    name_p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.0)
    _style_run(name_p.add_run(spec.get('name', '')), size=NAME_SIZE, bold=True,
               upper=True)

    title_p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.0)
    _style_run(title_p.add_run(spec.get('title', '')), size=TITLE_SIZE, bold=True)

    contact_p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3.0)
    _style_run(contact_p.add_run(spec.get('contact', '')), size=CONTACT_SIZE)


def _add_section_heading(doc, heading):
    p = _new_paragraph(doc, space_before=6.0, space_after=2.0)
    _style_run(p.add_run(heading), size=SECTION_SIZE, bold=True, upper=True)
    _add_bottom_border(p)


def build(spec: dict, output_path: str) -> str:
    doc = Document()
    _configure_document(doc)
    _add_header(doc, spec)

    for section in spec.get('sections', []):
        stype = section.get('type')
        renderer = SECTION_RENDERERS.get(stype)
        if renderer is None:
            raise ValueError(
                f"Unknown section type: {stype!r} "
                f"(expected one of {sorted(SECTION_RENDERERS)})"
            )
        _add_section_heading(doc, section.get('heading', ''))
        renderer(doc, section)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return str(out)


def main():
    if len(sys.argv) < 3:
        print("Usage: python build_docx_cv.py <spec.json> <output.docx>",
              file=sys.stderr)
        sys.exit(1)

    spec_path = Path(sys.argv[1])
    output_path = sys.argv[2]

    try:
        spec = json.loads(spec_path.read_text(encoding='utf-8'))
        result = build(spec, output_path)
        print(f"Successfully generated Word CV: {result}")
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
